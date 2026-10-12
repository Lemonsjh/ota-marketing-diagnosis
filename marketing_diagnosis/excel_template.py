from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from marketing_diagnosis.excel_template_blocks import _blocks as _legacy_blocks
from marketing_diagnosis.excel_template_current import update_blocks


TEMPLATE_FILE_NAME = "s14_ota_diagnosis_template.xlsx"

DARK = "155B4D"
GREEN = "1E8F67"
HEADER = "2B6F63"
LIGHT_GREEN = "E6F5EF"
SAMPLE = "F2FAF7"
INPUT = "FFF5CC"
ORANGE = "FFF0D8"
WHITE = "FFFFFF"
TEXT = "26343D"
MUTED = "667784"
BORDER = "D7E4DF"

YES_NO = ["是", "否"]
def _blocks():
    return update_blocks(_legacy_blocks())



def _style_sheet(sheet, subtitle: str) -> None:
    sheet.merge_cells("A1:L1")
    sheet["A1"] = sheet.title
    sheet["A1"].fill = PatternFill("solid", fgColor=DARK)
    sheet["A1"].font = Font(color=WHITE, bold=True, size=15)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells("A2:L2")
    sheet["A2"] = subtitle
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    sheet["A2"].font = Font(color=MUTED, size=10)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 36
    sheet.freeze_panes = "A3"


def _apply_validation(sheet, column: int, start: int, end: int, values: list[Any]) -> None:
    quoted = ",".join(str(value) for value in values)
    validation = DataValidation(type="list", formula1=f'"{quoted}"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"{get_column_letter(column)}{start}:{get_column_letter(column)}{end}")


def _write_block(sheet, start_row: int, block: dict[str, Any]) -> tuple[int, str]:
    headers = block["headers"]
    last_column = len(headers)
    last_letter = get_column_letter(last_column)
    border = Border(*(Side(style="thin", color=BORDER) for _ in range(4)))

    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=last_column)
    title_cell = sheet.cell(start_row, 1, f'{block["title"]} 〔系统模块：{block["module"]}〕')
    title_cell.fill = PatternFill("solid", fgColor=GREEN)
    title_cell.font = Font(color=WHITE, bold=True, size=11)
    title_cell.alignment = Alignment(vertical="center")

    sheet.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=last_column)
    note_cell = sheet.cell(start_row + 1, 1, block["note"])
    note_cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    note_cell.font = Font(color=MUTED, size=10)
    note_cell.alignment = Alignment(wrap_text=True, vertical="center")

    header_row = start_row + 2
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, value)
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[header_row].height = 30

    sample_start = header_row + 1
    for row_offset, values in enumerate(block["samples"]):
        row = sample_start + row_offset
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.fill = PatternFill("solid", fgColor=SAMPLE)
            cell.font = Font(color=TEXT, size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

    blank_start = sample_start + len(block["samples"])
    blank_end = blank_start + block["blank_rows"] - 1
    for row in range(blank_start, blank_end + 1):
        sheet.cell(row, 1, "否")
        sheet.cell(row, 2, block["module"])
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=INPUT)
            cell.font = Font(color=TEXT, size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

    data_end = max(sample_start + len(block["samples"]) - 1, blank_end)
    _apply_validation(sheet, 1, sample_start, data_end, YES_NO)
    sheet.conditional_formatting.add(
        f"A{sample_start}:A{data_end}",
        FormulaRule(formula=[f'A{sample_start}="是"'], fill=PatternFill("solid", fgColor="DFF4EA")),
    )

    header_index = {name: index + 1 for index, name in enumerate(headers)}
    for header, values in block["validations"].items():
        if header in header_index:
            _apply_validation(sheet, header_index[header], sample_start, data_end, values)

    for index, header in enumerate(headers, start=1):
        width = 15
        if header == "系统模块（勿改）":
            width = 35
        elif any(token in header for token in ("名称", "房型", "内容", "说明", "范围", "地址")):
            width = 25
        elif any(token in header for token in ("日期", "时间")):
            width = 19
        sheet.column_dimensions[get_column_letter(index)].width = width

    return data_end + 2, f"'{sheet.title}'!A{sample_start}:A{data_end}"


def build_excel_template(path: str | Path) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    guide = workbook.create_sheet("00_填写说明")
    guide.merge_cells("A1:F1")
    guide["A1"] = "S14 酒店诊断填报模板"
    guide["A1"].fill = PatternFill("solid", fgColor=DARK)
    guide["A1"].font = Font(color=WHITE, bold=True, size=15)
    guide["A1"].alignment = Alignment(vertical="center")
    guide.row_dimensions[1].height = 34
    guide.merge_cells("A2:F2")
    guide["A2"] = "绿色行为示例，黄色行为新增区域。示例行默认不导入；覆盖真实数据后，将是否导入改为“是”。工作表名称、系统模块列和表头不要修改。"
    guide["A2"].fill = PatternFill("solid", fgColor=ORANGE)
    guide["A2"].font = Font(color="8A4B00", bold=True, size=10)
    guide["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    guide.row_dimensions[2].height = 42
    rules = [
        ["序号", "填写规则", "正确做法", "错误示例", "是否必须", "影响"],
        [1, "示例数据", "替换为当前酒店ID和名称", "保留示例值直接上传", "必须", "酒店不匹配时拒绝导入"],
        [2, "是否导入", "真实行改为是，空行保持否", "全部空行改为是", "必须", "产生空数据"],
        [3, "比例", "填0.6129或61.29%", "填61.29但不带%", "必须", "比例放大100倍"],
        [4, "经营同比", "有去年同期时同时填写", "把缺失同期填成0", "有数据时", "缺同期则YOY不计算"],
        [5, "未知数据", "保持空白", "用0表示未知", "否", "可能误判经营差"],
        [6, "标注(%)的字段", "填10或10%表示10%", "把0.10当成10%", "对应字段", "0.10数值表示0.10%"],
    ]
    for row in rules:
        guide.append(row)
    for cell in guide[3]:
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in guide.iter_rows(min_row=4, max_row=guide.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = Border(*(Side(style="thin", color=BORDER) for _ in range(4)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in enumerate([8, 24, 38, 30, 12, 28], start=1):
        guide.column_dimensions[get_column_letter(column)].width = width
    guide.freeze_panes = "A3"

    blocks = _blocks()
    sheet_subtitles = {
        "02_基础经营": "酒店基础信息、日度经营、本期与去年同期经营数据。",
        "03_房型商品": "房型出租率、RevPAR，以及平台在售商品和价格。",
        "04_流量用户": "美团流量与携程流量、PSI分别填写。百分号字段(%)使用0–100数值。",
        "05_推广竞对": "同一近30天周期的推广投入、推广预订订单金额与竞对流失。",
        "06_口碑权益": "点评平台概览、权益中心和页面基础信息。",
        "07_配置视频挂冠": "固定配置状态、视频数量和酒店挂冠。",
        "08_可选补充": "扫码订单、周边活动和竞品价格；没有数据可以全部留空。",
    }
    progress: list[tuple[dict[str, Any], str]] = []
    for sheet_name, subtitle in sheet_subtitles.items():
        sheet = workbook.create_sheet(sheet_name)
        _style_sheet(sheet, subtitle)
        row = 4
        for block in [item for item in blocks if item["sheet"] == sheet_name]:
            row, import_range = _write_block(sheet, row, block)
            progress.append((block, import_range))

    dashboard = workbook.create_sheet("01_填报总览", 1)
    dashboard.merge_cells("A1:G1")
    dashboard["A1"] = "S14 客户填报进度与页面覆盖"
    dashboard["A1"].fill = PatternFill("solid", fgColor=DARK)
    dashboard["A1"].font = Font(color=WHITE, bold=True, size=15)
    dashboard.merge_cells("A2:G2")
    dashboard["A2"] = "下表统计选择导入的行数，不代表校验已通过。仅填写酒店实际使用的平台，缺失数据留空。"
    dashboard["A2"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    dashboard["A2"].font = Font(color=MUTED, size=10)
    dashboard["A2"].alignment = Alignment(wrap_text=True)
    dashboard.append([])
    dashboard.append(["序号", "页面/模块", "填写位置", "建议行数", "选择导入行数", "填报状态", "说明"])
    for cell in dashboard[4]:
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    required_by_module = {
        "basic_info": 1, "hotel_daily": 7, "hotel_performance_daily": 8,
        "room_type_performance_daily": 2, "products": 1, "exposure_daily": 7,
        "ota_funnel": 7, "user_source_monthly": 1, "promotion_finance": 1,
        "promotion_revenue": 1, "order_loss_monthly": 1, "review_overviews": 1,
        "joined_rights": 1, "promotion_status": 7, "video_upload_status": 3,
        "manual_inputs": 1, "scan_orders": 0, "nearby_events": 0, "competitors": 0,
    }
    for index, (block, import_range) in enumerate(progress, start=1):
        row = dashboard.max_row + 1
        required = required_by_module.get(block["module"], 1)
        if block["module"].startswith("ctrip_"):
            required = 0
        dashboard.append([index, block["title"], block["sheet"], required, None, None, block["page"]])
        dashboard.cell(row, 5, f'=COUNTIF({import_range},"是")')
        if required == 0:
            dashboard.cell(row, 6, f'=IF(E{row}>0,"已补充","可选")')
        else:
            dashboard.cell(row, 6, f'=IF(E{row}>={required},"已填写",IF(E{row}>0,"部分填写","未填写"))')
        for cell in dashboard[row]:
            cell.border = Border(*(Side(style="thin", color=BORDER) for _ in range(4)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in enumerate([8, 30, 22, 14, 14, 16, 34], start=1):
        dashboard.column_dimensions[get_column_letter(column)].width = width
    dashboard.freeze_panes = "A4"
    dashboard.row_dimensions[2].height = 36

    workbook.save(output_path)
    return output_path


def ensure_excel_template(output_root: str | Path) -> Path:
    template_path = Path(output_root) / "_templates" / TEMPLATE_FILE_NAME
    build_excel_template(template_path)
    return template_path
