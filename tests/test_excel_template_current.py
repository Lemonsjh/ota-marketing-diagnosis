from datetime import date
import pytest
from openpyxl import Workbook, load_workbook

from marketing_diagnosis.customer_excel_current import FIELDS, load_customer_excel_workbook
from marketing_diagnosis.customer_excel_result import enrich_customer_excel_result
from marketing_diagnosis.excel_loader import load_excel_dataset
from marketing_diagnosis.excel_template import _blocks, build_excel_template
from marketing_diagnosis.excel_fields import EXTRA_FIELD_MAP
from marketing_diagnosis.latest import normalize_dataset, process


def _block(sheet, module, values, *, title=True, explicit=True):
    if title:
        sheet.append([f'填写区域〔系统模块：{module}〕'])
    fields = FIELDS[module]
    sheet.append(['是否导入', '系统模块（勿改）', *fields])
    sheet.append(['是', module if explicit else None, *(values.get(key) for key in fields.values())])


def test_every_reserved_input_row_has_module_and_known_headers(tmp_path):
    path = build_excel_template(tmp_path / 'template.xlsx')
    book = load_workbook(path)
    for sheet in book:
        module = None
        for row in sheet:
            if isinstance(row[0].value, str) and '〔系统模块：' in row[0].value:
                module = row[0].value.split('〔系统模块：', 1)[1].split('〕', 1)[0]
            if row[0].value == '是否导入':
                labels = [cell.value for cell in row if cell.value]
                assert set(labels[2:]) <= set(FIELDS[module])
            elif row[0].value == '否' and module:
                assert row[1].value == module
    assert 'promotion_revenue' not in {block['module'] for block in _blocks()}
    assert load_excel_dataset(path)['__source_diagnostics__'][0]['imported_rows'] == 0
    book.close()


def test_complete_old_block_can_infer_blank_module_but_fragment_is_rejected():
    book = Workbook()
    _block(book.active, 'basic_info', {'hotel_id': 'audit', 'hotel_name': '测试酒店'}, explicit=False)
    assert load_customer_excel_workbook(book)['__excel_context__']['hotel_id'] == 'audit'
    fragment = Workbook()
    _block(fragment.active, 'basic_info', {'hotel_id': 'audit'}, title=False, explicit=False)
    with pytest.raises(ValueError, match='第2行'):
        load_customer_excel_workbook(fragment)


def test_selected_empty_or_incomplete_promotion_cannot_become_real_zero():
    for values in ({}, {'hotel_id': 'audit', 'spend_amount': 0}):
        book = Workbook()
        _block(book.active, 'promotion_finance', values)
        with pytest.raises(ValueError):
            load_customer_excel_workbook(book)


@pytest.mark.parametrize('invalid', ['未知', 'NaN', 'Infinity', -1])
def test_invalid_promotion_value_is_not_silently_zero(invalid):
    book = Workbook()
    _block(book.active, 'promotion_finance', {'hotel_id': 'audit', 'spend_amount': invalid,
           'booking_order_amount': 0, 'period_start_date': '2026-09-13', 'period_end_date': '2026-10-12'})
    with pytest.raises(ValueError):
        load_customer_excel_workbook(book)


def test_actual_upload_pipeline_uses_new_fields_and_independent_channels(tmp_path):
    book = Workbook()
    base = {'hotel_id': 'audit', 'snapshot_time': '2026-10-12 18:00:00'}
    period = {'period_start_date': '2026-09-13', 'period_end_date': '2026-10-12'}
    _block(book.active, 'basic_info', {'hotel_id': 'audit', 'hotel_name': '测试酒店'})
    _block(book.active, 'promotion_finance', {**base, **period, 'spend_amount': 2000, 'booking_order_amount': 24000})
    _block(book.active, 'review_overviews', {**base, 'review_platform': '美团', 'review_score': 4.8,
           'total_review_count': 100, 'unreplied_review_count': 0, 'yesterday_new_review_count': 0, 'yesterday_review_date': date(2026, 10, 11)})
    _block(book.active, 'ctrip_promotion_status', {**base, 'activity_code': '信息完整度', 'metric_value': 100, 'status': '已开通'})
    _block(book.active, 'ctrip_promotion_status', {**base, 'activity_code': '积分联盟', 'status': '未报名'})
    _block(book.active, 'ctrip_promotion_status', {**base, 'activity_code': '闪住', 'status': '已参加'})
    _block(book.active, 'ctrip_promotion_activity', {**base, 'activity_name': 'YOYO卡&扫码住专享促销', 'status': 'JOINED'})
    _block(book.active, 'ctrip_promotion_activity_performance', {**base, 'campaign_name': 'YOYO卡&扫码住专享促销', 'campaign_quantity': 150})
    _block(book.active, 'ctrip_video_upload_status', {**base, 'video_type': '详情页视频', 'uploaded_count': 0, 'required_count': 1})
    _block(book.active, 'ctrip_psi_score', {**base, 'psi_total_score': 5.6})
    _block(book.active, 'ctrip_review_overview', {**base, 'channel_source': '携程', 'review_score': 4.8, 'total_review_count': 80, 'unreplied_review_count': 0})
    _block(book.active, 'ctrip_review_yesterday', {**base, 'platform_scope': '携程', 'yesterday_new_review_count': 2, 'yesterday_review_date': '2026-10-11'})
    path = tmp_path / 'upload.xlsx'
    book.save(path)
    raw = load_excel_dataset(path)
    normalized = normalize_dataset(raw)
    result = enrich_customer_excel_result(process(normalized), raw)
    items = {item['standard_item_id']: item for item in result['visual_diagnosis']['items']}
    assert items[9]['promotion_performance']['roi'] == 12
    assert items[9]['item_score'] == 8
    assert items[13]['yesterday_new_review_count'] == 0
    assert items[13]['yesterday_review_date'] == '2026-10-11'
    assert '数据库' not in next(field for field in items[13]['fields'] if field['label'] == '昨日新增点评数')['origin']
    ctrip = result['ctrip_items']
    assert ctrip['8']['item_score'] == 4
    assert ctrip['6']['item_score'] == 8
    assert ctrip['20']['fields'][0]['value'] == '0/1'
    assert ctrip['20']['item_score'] == 0
    assert ctrip['14']['item_score'] == 0
    assert ctrip['17']['item_score'] == 2
    assert ctrip['7']['item_score'] == 5
    assert ctrip['12']['platforms'][0]['yesterday_new_review_count'] == 2


def test_percent_label_accepts_number_text_and_excel_percent():
    for value, number_format in [(10, 'General'), ('10%', 'General'), (.1, '0%')]:
        book = Workbook()
        _block(book.active, 'ctrip_business_metrics_funnel', {'hotel_id': 'audit', 'exposure_to_detail_rate_pct': value})
        labels = [cell.value for cell in book.active[2]]
        book.active.cell(3, labels.index('曝光到详情转化率(%)') + 1).number_format = number_format
        assert load_customer_excel_workbook(book)['ctrip_business_metrics_funnel'][0]['exposure_to_detail_rate_pct'] == 10


def test_upload_cannot_override_trusted_hotel():
    from importlib import import_module
    runtime = import_module('skills.s14-operation-diagnosis.runtime')
    with pytest.raises(ValueError, match='当前酒店不一致'):
        runtime._apply_excel_context({'hotel_id': 'bound_hotel'}, {
            '__excel_context__': {'hotel_id': 'other_hotel'},
        })
    with pytest.raises(ValueError, match='当前酒店不一致'):
        runtime._apply_excel_context({'hotel_id': 'bound_hotel'}, {
            '__excel_context__': {}, 'hotel_daily': [{'hotel_id': 'other_hotel'}],
        })


def test_workbook_must_not_mix_hotel_rows():
    book = Workbook()
    _block(book.active, 'basic_info', {'hotel_id': 'one_hotel'})
    _block(book.active, 'hotel_daily', {'hotel_id': 'other_hotel', 'room_count': 20})
    with pytest.raises(ValueError, match='同一家酒店'):
        load_customer_excel_workbook(book)


@pytest.mark.parametrize('section', list(EXTRA_FIELD_MAP))
def test_all_ctrip_sections_survive_actual_normalization(section):
    book = Workbook()
    values = {key: 1 for key in FIELDS[section].values()}
    values.update(hotel_id='audit', period_start_date='2026-09-13', period_end_date='2026-10-12',
                  snapshot_time='2026-10-12 18:00:00', business_date='2026-10-12', yesterday_review_date='2026-10-11')
    _block(book.active, section, values)
    raw = load_customer_excel_workbook(book)
    assert normalize_dataset(raw)['sections'][section] == raw[section]
